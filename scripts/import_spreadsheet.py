"""Imports the old weekly spreadsheet into the site.

    python scripts/import_spreadsheet.py reference/Week04_2026.xlsx

- WEEK tab: that week's spreads, records and everyone's picks (games and
  results themselves come from ESPN).
- SEASON tab: weekly point totals for earlier weeks (any week that isn't
  the WEEK tab's week), stored as "imported".

Needs openpyxl. Reads the private admin key from apps-script/Private.js.
"""

import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent


def admin_call(task, payload):
    key = re.search(r"adminKey: '([0-9a-f]+)'", (ROOT / "apps-script/Private.js").read_text()).group(1)
    url = re.search(r"apiUrl: '([^']+)'", (ROOT / "site/config.js").read_text()).group(1)
    body = json.dumps({"action": "adminTask", "key": key, "task": task, **payload})
    # curl rather than urllib: some Python installs on macOS lack SSL certificates.
    out = subprocess.run(
        ["curl", "-sL", "--max-time", "300", "-H", "Content-Type: text/plain;charset=utf-8", "--data-binary", "@-", url],
        input=body, capture_output=True, text=True, check=True,
    ).stdout
    result = json.loads(out)
    if not result.get("ok"):
        sys.exit(f"{task} failed: {result.get('error')}")
    return result["data"]


def read_week(ws):
    week = int(re.search(r"WEEK #(\d+)", str(ws["A1"].value)).group(1))
    spread_source = "FanDuel" if "fandu" in str(ws["B2"].value).lower() else ""

    names = {}
    for col in range(6, ws.max_column + 1):
        name = ws.cell(3, col).value
        if name and str(name).strip():
            names[col] = str(name).strip()

    games, picks = [], []
    row = 4
    while ws.cell(row, 1).value and ws.cell(row + 1, 1).value:
        pair = [row, row + 1]
        teams = [str(ws.cell(r, 1).value).strip() for r in pair]
        favorite, spread = "", None
        for r in pair:
            value = ws.cell(r, 2).value
            if isinstance(value, (int, float)):
                favorite, spread = str(ws.cell(r, 1).value).strip(), abs(value)
        records = {}
        for r in pair:
            wins, losses = ws.cell(r, 3).value, ws.cell(r, 4).value
            if isinstance(wins, (int, float)) and isinstance(losses, (int, float)):
                records[str(ws.cell(r, 1).value).strip()] = f"{int(wins)}-{int(losses)}"
        games.append({"teams": teams, "favorite": favorite, "spread": spread, "records": records})
        for r in pair:
            for col, name in names.items():
                points = ws.cell(r, col).value
                if isinstance(points, (int, float)):
                    picks.append({"name": name, "team": str(ws.cell(r, 1).value).strip(), "points": int(points)})
        row += 2
    return week, spread_source, games, picks


def read_season(ws, skip_week):
    names = {col: str(ws.cell(1, col).value).strip() for col in range(2, ws.max_column + 1) if ws.cell(1, col).value}
    totals = []
    for row in range(2, ws.max_row + 1):
        week = ws.cell(row, 1).value
        if not isinstance(week, int) or week == skip_week:
            continue
        for col, name in names.items():
            points = ws.cell(row, col).value
            if isinstance(points, (int, float)):
                totals.append({"week": week, "name": name, "points": int(points)})
    return totals


def main():
    wb = openpyxl.load_workbook(sys.argv[1], data_only=False)
    week, spread_source, games, picks = read_week(wb["WEEK"])
    totals = read_season(wb["SEASON"], skip_week=week)

    per_player = {}
    for p in picks:
        per_player.setdefault(p["name"], []).append(p["points"])
    n = len(games)
    bad = {name: sorted(pts) for name, pts in per_player.items() if sorted(pts) != list(range(1, n + 1))}
    print(f"Week {week}: {n} games, {len(picks)} picks from {len(per_player)} players" + (f", PROBLEMS: {bad}" if bad else ", all valid"))
    print(f"SEASON: {len(totals)} weekly totals for weeks {sorted({t['week'] for t in totals})}")

    print("importTotals:", admin_call("importTotals", {"totals": totals}))
    print("importWeek:", admin_call("importWeek", {"week": week, "spreadSource": spread_source, "games": games, "picks": picks}))


if __name__ == "__main__":
    main()
